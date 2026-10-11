# -*- coding: utf-8 -*-
"""Markdown to DOCX converter with Persian RTL support and academic formatting."""

import os
import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

PERSIAN_BODY_FONT = "B Nazanin"
PERSIAN_HEADING_FONT = "B Titr"
FALLBACK_FONT = "Tahoma"
CODE_FONT = "Consolas"

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_paragraph_rtl(paragraph):
    """Set Right-to-Left bidirectional property on paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def apply_font(run, font_name=PERSIAN_BODY_FONT, size_pt=12, bold=False, italic=False, color_rgb=None):
    """Apply font and complex script fonts to run."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

def convert_md_to_docx(md_path="Thesis.md", docx_path="Thesis.docx"):
    md_file = Path(md_path).resolve()
    if not md_file.is_file():
        raise FileNotFoundError(f"Markdown file not found: {md_file}")

    with open(md_file, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # Configure A4 page layout with 2.54 cm (1 inch) margins
    for section in doc.sections:
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    lines = content.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal in_code_block, code_lines
        if code_lines:
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)

            # Code box background
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
            pPr.append(shd)

            run = p.add_run(code_text)
            apply_font(run, font_name=CODE_FONT, size_pt=9.5, bold=False, color_rgb=RGBColor(40, 44, 52))
            code_lines = []
        in_code_block = False

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        rows_data = []
        for line in table_lines:
            # Strip outer pipes
            trimmed = line.strip()
            if trimmed.startswith("|"):
                trimmed = trimmed[1:]
            if trimmed.endswith("|"):
                trimmed = trimmed[:-1]
            cells = [c.strip() for c in trimmed.split("|")]
            # Filter divider rows like :--- | :---:
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            rows_data.append(cells)

        if rows_data:
            num_rows = len(rows_data)
            num_cols = max(len(r) for r in rows_data)

            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            # Set table grid border
            tblPr = table._tbl.tblPr
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
                f'<w:left w:val="none"/>'
                f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="888888"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr.append(borders)

            for r_idx, row in enumerate(rows_data):
                is_header = (r_idx == 0)
                for c_idx in range(num_cols):
                    cell = table.cell(r_idx, c_idx)
                    cell_text = row[c_idx] if c_idx < len(row) else ""
                    # Background
                    if is_header:
                        set_cell_background(cell, "2B4C7E")
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, "F9FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")

                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

                    p = cell.paragraphs[0]
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    set_paragraph_rtl(p)

                    run = p.add_run(cell_text)
                    if is_header:
                        apply_font(run, font_name=PERSIAN_HEADING_FONT, size_pt=11, bold=True, color_rgb=RGBColor(255, 255, 255))
                    else:
                        apply_font(run, font_name=PERSIAN_BODY_FONT, size_pt=10.5, bold=False, color_rgb=RGBColor(30, 30, 30))

            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(4)
            p_after.paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        # Code block toggle
        if stripped.startswith("```"):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Table detection
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table()

        # Horizontal rule / divider
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("―" * 30)
            apply_font(run, font_name=FALLBACK_FONT, size_pt=10, color_rgb=RGBColor(180, 180, 180))
            continue

        # Headings
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)
            run = p.add_run(title_text)
            apply_font(run, font_name=PERSIAN_HEADING_FONT, size_pt=18, bold=True, color_rgb=RGBColor(24, 43, 73))
            continue

        if stripped.startswith("## "):
            title_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)
            run = p.add_run(title_text)
            apply_font(run, font_name=PERSIAN_HEADING_FONT, size_pt=15, bold=True, color_rgb=RGBColor(31, 78, 121))
            continue

        if stripped.startswith("### "):
            title_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)
            run = p.add_run(title_text)
            apply_font(run, font_name=PERSIAN_HEADING_FONT, size_pt=13, bold=True, color_rgb=RGBColor(44, 98, 150))
            continue

        if stripped.startswith("#### "):
            title_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)
            run = p.add_run(title_text)
            apply_font(run, font_name=PERSIAN_HEADING_FONT, size_pt=12, bold=True, color_rgb=RGBColor(60, 60, 60))
            continue

        # Blockquote (Callouts)
        if stripped.startswith(">"):
            quote_text = stripped[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)

            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
            pPr.append(shd)

            run = p.add_run(quote_text)
            apply_font(run, font_name=PERSIAN_BODY_FONT, size_pt=11.5, italic=True, color_rgb=RGBColor(40, 60, 90))
            continue

        # Bullet point or list item
        if stripped.startswith(("- ", "* ", "• ")):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)

            bullet_run = p.add_run("• ")
            apply_font(bullet_run, font_name=FALLBACK_FONT, size_pt=12, bold=True, color_rgb=RGBColor(31, 78, 121))

            text_run = p.add_run(bullet_text)
            apply_font(text_run, font_name=PERSIAN_BODY_FONT, size_pt=12, bold=False, color_rgb=RGBColor(20, 20, 20))
            continue

        # Regular paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)

            run = p.add_run(stripped)
            apply_font(run, font_name=PERSIAN_BODY_FONT, size_pt=12, bold=False, color_rgb=RGBColor(20, 20, 20))

    # Flush any remaining items
    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    out_docx = Path(docx_path).resolve()
    doc.save(str(out_docx))
    print("=" * 60)
    print(f"DOCX successfully exported: {out_docx}")
    print(f"File size: {out_docx.stat().st_size:,} bytes")
    print("=" * 60)

if __name__ == "__main__":
    convert_md_to_docx()
